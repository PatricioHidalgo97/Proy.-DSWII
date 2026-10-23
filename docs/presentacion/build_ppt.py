#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Genera Presentacion_Microservicios_Grupo2.pptx para el cierre del proyecto
"Sistema de Gestión Administrativa y de Caja" (del monolito a microservicios).
Diseño con la paleta del frontend (src/styles.css):
  navy #10253f, navy_soft #1a3659, blue #1769e0, ink #172033,
  muted #657089, line #dfe5ee, bg #f3f6fa.
Requisito: pip install python-pptx  (ya instalado en este entorno).
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ---------------------------------------------------------------- paleta
NAVY      = RGBColor(0x10, 0x25, 0x3F)
NAVY_SOFT = RGBColor(0x1A, 0x36, 0x59)
BLUE      = RGBColor(0x17, 0x69, 0xE0)
BLUE_DARK = RGBColor(0x0F, 0x54, 0xBB)
BLUE_LIGHT= RGBColor(0xE8, 0xF0, 0xFD)
INK       = RGBColor(0x17, 0x20, 0x33)
MUTED     = RGBColor(0x65, 0x70, 0x89)
LINE      = RGBColor(0xDF, 0xE5, 0xEE)
BG        = RGBColor(0xF3, 0xF6, 0xFA)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
GREEN     = RGBColor(0x18, 0x79, 0x4E)
GREEN_L   = RGBColor(0xEA, 0xF7, 0xEF)
ORANGE    = RGBColor(0xDC, 0x77, 0x26)
ORANGE_L  = RGBColor(0xFF, 0xF2, 0xDC)
VIOLET    = RGBColor(0x76, 0x57, 0xD6)
VIOLET_L  = RGBColor(0xEF, 0xE9, 0xFF)
RED       = RGBColor(0xC4, 0x3D, 0x3D)
RED_L     = RGBColor(0xFF, 0xF0, 0xF0)
FONT      = "Segoe UI"

EMU_W, EMU_H = Inches(13.333), Inches(7.5)

prs = Presentation()
prs.slide_width = EMU_W
prs.slide_height = EMU_H
BLANK = prs.slide_layouts[6]

def add_slide(bg=BG):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = bg
    return s

def rect(slide, x, y, w, h, fill=None, line=None, line_w=0.75, shape=MSO_SHAPE.RECTANGLE, shadow=False, adjust=None):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(line_w)
    sp.shadow.inherit = False
    if adjust is not None:
        try: sp.adjustments[0] = adjust
        except Exception: pass
    return sp

def tb(slide, x, y, w, h, text="", size=13, color=INK, bold=False, font=FONT,
       align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=1.15,
       space_after=0, caps=False, italic=False, autosize=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    lines = text.split("\n")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        p.space_after = Pt(space_after)
        r = p.add_run(); r.text = ln.upper() if caps else ln
        r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
        r.font.name = font; r.font.color.rgb = color
    return box

def bullets(slide, x, y, w, h, items, size=12.5, color=INK, gap=8,
            bullet_color=BLUE, lead_bold=True, line_spacing=1.18):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.word_wrap = True
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    first = True
    for item in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.line_spacing = line_spacing
        p.space_after = Pt(gap)
        rb = p.add_run(); rb.text = "▸  "; rb.font.size = Pt(size)
        rb.font.bold = True; rb.font.name = FONT; rb.font.color.rgb = bullet_color
        if "**" in item:
            lead, rest = item.split("**", 1)
            r1 = p.add_run(); r1.text = lead
            r1.font.size = Pt(size); r1.font.name = FONT
            r1.font.bold = lead_bold; r1.font.color.rgb = color
            r2 = p.add_run(); r2.text = rest
            r2.font.size = Pt(size); r2.font.name = FONT
            r2.font.bold = False; r2.font.color.rgb = MUTED
        else:
            r = p.add_run(); r.text = item
            r.font.size = Pt(size); r.font.name = FONT
            r.font.color.rgb = color
    return box

def card(slide, x, y, w, h, accent=BLUE, title=None, items=None, body=None,
         size=12.5, fill=WHITE, title_size=13.5, title_color=NAVY,
         body_color=INK, radius=0.045):
    rect(slide, x, y, w, h, fill=fill, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adjust=radius)
    rect(slide, x, y + 0.18, 0.07, h - 0.36, fill=accent, shape=MSO_SHAPE.RECTANGLE)
    cy = y + 0.22
    if title:
        tb(slide, x + 0.30, cy, w - 0.62, 0.42, title, size=title_size,
           bold=True, color=title_color)
        cy += 0.44
    if items:
        bullets(slide, x + 0.30, cy, w - 0.60, h - cy + y - 0.10, items,
                size=size, bullet_color=accent, gap=7)
    if body:
        tb(slide, x + 0.30, cy if not items else cy + len(items) * 0.30,
           w - 0.60, h - cy + y - 0.10, body, size=size, color=body_color)
    return card

def kpi(slide, x, y, w, h, value, label, accent=BLUE, vsize=24, lsize=10.5,
        fill=WHITE, vcolor=NAVY):
    rect(slide, x, y, w, h, fill=fill, line=LINE,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, adjust=0.10)
    rect(slide, x + w - 0.12, y + 0.18, 0.06, h - 0.36, fill=accent,
         shape=MSO_SHAPE.RECTANGLE)
    tb(slide, x + 0.24, y + 0.14, w - 0.45, 0.55, value, size=vsize, bold=True,
       color=vcolor)
    tb(slide, x + 0.24, y + h - 0.34, w - 0.45, 0.30, label, size=lsize,
       color=MUTED, bold=True)

def picture(slide, path, x, y, w, h, border=True, radius=None):
    pic = slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    if border:
        rect(slide, x - 0.015, y - 0.015, w + 0.03, h + 0.03, fill=None,
             line=LINE, line_w=1.0,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE if radius is None else MSO_SHAPE.ROUNDED_RECTANGLE,
             adjust=radius)
    return pic

def arrow_right(slide, x, y, w=0.42, h=0.20, color=MUTED):
    return rect(slide, x, y, w, h, fill=color, shape=MSO_SHAPE.RIGHT_ARROW)

def footer(slide, n, total=None):
    rect(slide, 0.75, 6.96, 11.833, 0.012, fill=LINE)
    tb(slide, 0.75, 7.04, 9.5, 0.3,
       "Grupo 2 · Sistema de Gestión Administrativa y de Caja — De monolito a microservicios",
       size=8.5, color=MUTED)
    total_str = f"{n}" if total is None else f"{n} / {total}"
    tb(slide, 12.35, 7.04, 0.25, 0.3, total_str, size=8.5, color=MUTED,
       align=PP_ALIGN.RIGHT)

def header(slide, eyebrow, title, sub=None, title_size=31, sub_y=1.42):
    rect(slide, 0.75, 0.42, 0.55, 0.055, fill=BLUE)
    tb(slide, 0.75, 0.52, 11.833, 0.30, eyebrow, size=10.5, color=BLUE,
       bold=True, caps=True)
    tb(slide, 0.75, 0.78, 11.833, 0.62, title, size=title_size, color=NAVY, bold=True)
    if sub:
        tb(slide, 0.75, sub_y, 11.833, 0.40, sub, size=12, color=MUTED)

def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text

IM = "docs/presentacion/capturas"

# ================================================================ SLIDE 1 · CARÁTULA
s = add_slide(bg=NAVY)
rect(s, 0, 0, 13.333, 7.5, fill=NAVY)
rect(s, 0, 0, 13.333, 0.16, fill=BLUE)
rect(s, 0, 7.34, 13.333, 0.16, fill=BLUE)
# decoración derecha
rect(s, 9.6, 0.6, 3.733, 6.3, fill=NAVY_SOFT)
rect(s, 10.05, 1.05, 2.85, 0.10, fill=BLUE)
tb(s, 10.05, 1.30, 2.85, 0.35, "GRUPO 2", size=13, color=BLUE_DARK, bold=True)

tb(s, 0.9, 2.05, 8.4, 1.7, "Proyecto de cierre", size=13, color=RGBColor(0x9F, 0xC0, 0xFF), bold=True)
tb(s, 0.9, 2.42, 8.5, 1.9, "Del monolito a microservicios", size=44, color=WHITE, bold=True)
tb(s, 0.9, 3.62, 8.6, 0.85, "Sistema de Gestión Administrativa y de Caja — arquitectura dividida en identidad, catálogo y caja detrás de un API Gateway.", size=15, color=RGBColor(0xCF, 0xDD, 0xEA))
rect(s, 0.92, 4.62, 1.6, 0.045, fill=BLUE)

# nombres de los 5 integrantes
nombres = [
    ("1.", "Hidalgo Salas, José Patricio"),
    ("2.", "Calizaya Alvarez, Gerardo Matías"),
    ("3.", "Pajuelo Chipana, Juan Alonso Esteban"),
    ("4.", "Tupa Churata, Salvador Roque"),
    ("5.", "Carrillo Bautista, Oscar Eduardo"),
]
y = 5.15
for num, nom in nombres:
    tb(s, 0.9, y, 0.45, 0.4, num, size=12, color=BLUE, bold=True)
    tb(s, 1.38, y, 7.4, 0.4, nom, size=13.5, color=WHITE)
    y += 0.46

tb(s, 0.9, 6.85, 8.5, 0.4, "Curso: Desarrollo de Servicios Web II (DSW II) · Ciclo 6 · Lima, Perú · 2026",
   size=11, color=RGBColor(0xAE, 0xC1, 0xD6))
tb(s, 10.05, 6.85, 2.9, 0.4, "Exposición grupal · ~10 min", size=11,
   color=RGBColor(0xAE, 0xC1, 0xD6), align=PP_ALIGN.RIGHT)

notes(s, "Carátula (20 s). Saludar, decir nombres del grupo 2 y el título. Presentar brevemente de qué trata el proyecto: dividimos el monolito de gestión en microservicios sin cambiar la lógica ni el contrato.")

# ================================================================ SLIDE 2 · EL CASO (LORE)
s = add_slide()
header(s, "El caso", "Un mercado que cobra, y necesita orden",
       sub="Nació como un monolito académico para una asociación de comerciantes; el profesor pidió llevarlo a microservicios con API Gateway.")

card(s, 0.75, 1.95, 5.9, 4.55, accent=BLUE, title="Qué hace el sistema",
     items=[
         "**Socios:** registro, edición y consulta de asociados.",
         "**Giros, puestos y servicios:** cada puesto tiene un giro y servicios cobrables (fijos o por consumo).",
         "**Cobranza:** se generan cuentas individuales y masivas; se pagan o exoneran deudas en lote.",
         "**Caja:** recibos correlativos, canjes bancarios simulados e ingresos externos.",
         "**Tesorería:** egresos individuales o por CSV y reportes XLSX.",
     ], size=12.5)

card(s, 6.85, 1.95, 5.73, 2.30, accent=GREEN, title="El ciclo del negocio",
     items=[
         "Catálogos → Deudas → Cobranza y caja → Reportes",
     ], size=12.5)
# ciclo visual
ciclo = ["Catálogos", "Deudas", "Cobranza", "Reportes"]
cx = 7.15
for i, nombre in enumerate(ciclo):
    rect(s, cx, 4.42, 1.12, 0.52, fill=BLUE_LIGHT, line=BLUE,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, adjust=0.25)
    tb(s, cx, 4.51, 1.12, 0.34, nombre, size=10.5, bold=True, color=NAVY,
       align=PP_ALIGN.CENTER)
    if i < 3:
        arrow_right(s, cx + 1.18, 4.545, w=0.26, h=0.14, color=MUTED)
        cx += 1.48

card(s, 6.85, 4.38, 5.73, 2.12, accent=NAVY, title="El reto del grupo",
     items=[
         "Mismo contrato HTTP, mismas validaciones y permisos.",
         "Servicios independientes, cada uno con su base de datos.",
         "Sin reescribir el frontend Angular.",
     ], size=12.5)

notes(s, "Caso y lore (35 s). Explicar que el sistema ya existía como monolito y que el profesor pidió pasarlo a microservicios, manteniendo todo igual hacia afuera.")

# ================================================================ SLIDE 3 · EL PROBLEMA
s = add_slide()
header(s, "El problema", "Un monolito: cómodo al inicio, pesado al crecer",
       sub="Todo funcionaba, pero acoplado, frágil y difícil de repartir entre 5 personas.")

card(s, 0.75, 1.95, 5.75, 4.45, accent=RED, title="Dolores del monolito",
     items=[
         "**Un solo despliegue:** si falla la caja, cae todo el sistema.",
         "**Una base para todo:** catálogo y dinero están acoplados.",
         "**Cambios pequeños** obligan a recompilar y desplegar todo.",
         "**Difícil de dividir:** 5 personas tocando el mismo código.",
     ], size=12.5)

card(s, 6.85, 1.95, 5.73, 4.45, accent=GREEN, title="La meta del proyecto",
     items=[
         "**Servicios por dominio:** identidad, catálogo y caja, cada uno independiente.",
         "**Cada servicio, su base de datos:** sin claves foráneas cruzadas.",
         "**API Gateway:** una sola puerta de entrada para el frontend.",
         "**El monolito queda como respaldo** (9090), sin recibir tráfico del frontend.",
     ], size=12.5)

tb(s, 0.75, 6.52, 11.83, 0.4,
   "Regla de oro del proyecto: no cambiar la lógica de negocio ni el formato del JSON que consume el frontend.",
   size=12, color=MUTED, italic=True)

notes(s, "Problema y meta (40 s). Que quede claro por qué dividimos y cuáles son los límites: mismos endpoints, mismos permisos, monolito como respaldo.")

# ================================================================ SLIDE 4 · ARQUITECTURA FINAL
s = add_slide()
header(s, "Arquitectura final", "Angular → Gateway → 3 microservicios",
       sub="Estado 3 del patrón strangler: toda la API pública vive ya en microservicios.")

# tarjetas de servicios
def svc_card(x, y, w, h, nombre, puerto, rutas, bd, fill=WHITE, ac=BLUE):
    card(s, x, y, w, h, accent=ac, title=nombre, title_size=12.5, size=10.5)
    tb(s, x + 0.32, y + 0.55, w - 0.5, 0.3, f"puerto {puerto}", size=9.5, color=MUTED, bold=True)
    tb(s, x + 0.32, y + 0.86, w - 0.5, 0.62, rutas, size=9.5, color=INK)
    tb(s, x + 0.32, y + h - 0.30, w - 0.5, 0.26, bd, size=9, color=BLUE_DARK, bold=True)

svc_card(0.75, 1.95, 2.62, 1.85, "Angular", ":4400",
         "13 pantallas · 10 servicios HTTP · interceptor JWT",
         "Frontend SPA", fill=NAVY, ac=BLUE)
tb(s, 0.92, 2.20, 2.3, 0.35, "Angular", size=15, bold=True, color=WHITE)
tb(s, 0.92, 2.52, 2.3, 0.3, "puerto 4400", size=9.5, color=RGBColor(0x9F,0xC0,0xFF), bold=True)
tb(s, 0.92, 2.86, 2.3, 0.6, "13 pantallas · 10 servicios HTTP\ninterceptor JWT", size=9.5, color=RGBColor(0xCF,0xDD,0xEA))
tb(s, 0.92, 3.55, 2.3, 0.26, "Frontend SPA", size=9, color=BLUE_DARK, bold=True)

svc_card(4.05, 1.95, 2.42, 1.85, "API Gateway", ":9101",
         "10 rutas · /api/auth, socios, cash…\nprefijo desconocido → 404",
         "Sin Eureka", fill=BLUE, ac=BLUE_DARK)

svc_card(7.10, 1.95, 3.20, 1.25, "identity-service", ":9102",
         "/api/auth — login y token",
         "gestion_auth", ac=VIOLET)
svc_card(7.10, 3.35, 3.20, 1.25, "catalog-service", ":9103",
         "socios, giros, puestos, servicios, bancos",
         "gestion_catalog", ac=GREEN)
svc_card(7.10, 4.75, 3.20, 1.25, "cash-service", ":9104",
         "cuentas, recibos, egresos, reportes",
         "gestion_cash", ac=ORANGE)

# conectores gateway → servicios
bus_x = 6.62
rect(s, bus_x, 2.55, 0.04, 4.20, fill=LINE)
for yy, h in [(2.57, 1.25), (3.97, 1.25), (5.37, 1.25)]:
    arrow_right(s, bus_x + 0.05, yy + h/2 - 0.09, w=0.38, h=0.18, color=MUTED)
rect(s, 2.95, 2.55, 0.04, 1.85, fill=LINE)
arrow_right(s, 2.95 + 0.05, 3.47 - 0.09, w=0.42, h=0.18, color=MUTED)

tb(s, 0.75, 6.35, 11.83, 0.55,
   "Flujo: navegador (:4400) → gateway (:9101) → identity / catalog / cash. "
   "El monolito en :9090 sigue levantado como respaldo manual, pero el frontend ya no lo llama.",
   size=11, color=MUTED)

notes(s, "Arquitectura (50 s). Recorrer la flecha: Angular habla solo con el gateway, y el gateway reparte entre identity, catalog y cash. Cada servicio tiene su propia base de datos. Destacar que no hay Eureka, ni Kafka, ni contenedores: decisión de simplicidad.")

# ================================================================ SLIDE 5 · STACK
s = add_slide()
header(s, "Stack tecnológico", "¿Con qué lo construimos y por qué?",
       sub="Tecnologías elegidas por el equipo, cada una con su justificación.")

cols = [3.05, 4.15, 4.63]
rows_data = [
    ("Angular 21 + TypeScript", "Frontend SPA, 13 pantallas", "Ya existía; reescribir solo agregaba riesgo. Un cambio de línea apuntó al gateway."),
    ("Spring Boot 3.5 (Java 21)", "Base de los 4 servicios", "El equipo ya lo dominaba; ecosistema maduro y probado."),
    ("Spring Cloud Gateway", "Punto de entrada único", "Ruteo por prefijo sin servidor de descubrimiento (sin Eureka)."),
    ("Spring Security + JWT + BCrypt", "Autenticación y roles ADMIN/OPERADOR", "Mismo JWT en los 4 servicios; contraseñas nunca en claro."),
    ("Spring Data JPA + MySQL 8", "Persistencia relacional", "Una base de datos por servicio; consultas propias del negocio."),
    ("Flyway", "Migraciones versionadas", "Las 3 bases se crean desde cero con scripts, sin tocar a mano."),
    ("springdoc OpenAPI + Apache POI", "Contrato visible y reportes XLSX", "Swagger por servicio y exportación a Excel para el cliente."),
]
y = 1.85
rect(s, 0.75, y, sum(cols), 0.42, fill=NAVY)
for i, c in enumerate(cols):
    tb(s, 0.75 + sum(cols[:i]), y + 0.075, c, 0.3,
       ["Tecnología", "Para qué", "Por qué"][i], size=11, color=WHITE, bold=True)
y += 0.42
row_h = 0.52
for i, (tec, para, porq) in enumerate(rows_data):
    fill = WHITE if i % 2 == 0 else BLUE_LIGHT
    rect(s, 0.75, y, sum(cols), row_h, fill=fill, line=LINE, line_w=0.5)
    tb(s, 0.85, y + 0.09, cols[0] - 0.2, row_h - 0.16, tec, size=10.5, color=NAVY, bold=True)
    tb(s, 0.75 + cols[0] + 0.1, y + 0.09, cols[1] - 0.2, row_h - 0.16, para, size=10, color=INK)
    tb(s, 0.75 + cols[0] + cols[1] + 0.1, y + 0.09, cols[2] - 0.2, row_h - 0.16, porq, size=10, color=MUTED)
    y += row_h

tb(s, 0.75, 6.55, 11.83, 0.4,
   "Nota: todo se configura por variables de entorno (puertos, credenciales, URLs, JWT_SECRET); ningún valor se fija en el código.",
   size=11, color=MUTED, italic=True)

notes(s, "Stack (50 s). No leer toda la tabla: destacar Angular sin reescribir, Spring Boot/JWT/MySQL que ya conocían, Flyway para bases reproducibles, y por qué no pusimos Eureka/Kafka/Docker (simplicidad operativa).")

# ================================================================ SLIDE 6 · RECORRIDO
s = add_slide()
header(s, "Recorrido de la exposición", "Cinco bloques, uno por integrante",
       sub="~10 minutos en total; cada persona lleva su parte y conoce el flujo general.")

pasos = [
    ("1", "Entrada y seguridad", "Patricio", "Login, JWT y gateway", BLUE),
    ("2", "Datos maestros", "Gerardo", "Catálogos y fronteras de datos", GREEN),
    ("3", "Dinero y cobranza", "Juan", "Cuentas, recibos y correlativos", ORANGE),
    ("4", "Vista y reportes", "Salvador", "Tesorería, egresos y frontend", VIOLET),
    ("5", "Migración y calidad", "Oscar", "Strangler, Swagger, Postman, Flyway", NAVY),
]
y = 1.95
h = 0.82
for num, titulo, quien, detalle, ac in pasos:
    rect(s, 0.75, y, 11.833, h, fill=WHITE, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adjust=0.12)
    rect(s, 0.75, y, 0.07, h, fill=ac)
    tb(s, 1.0, y + h/2 - 0.13, 0.5, 0.3, num, size=15, color=ac, bold=True)
    tb(s, 1.62, y + 0.10, 3.1, 0.35, titulo, size=13, color=NAVY, bold=True)
    tb(s, 1.62, y + 0.45, 3.6, 0.3, detalle, size=10.5, color=MUTED)
    tb(s, 8.6, y + h/2 - 0.14, 2.6, 0.3, quien, size=12, color=ac, bold=True, align=PP_ALIGN.RIGHT)
    y += h + 0.14

tb(s, 0.75, 6.55, 11.83, 0.35,
   "Orden sugerido: carátula → caso → problema → arquitectura → stack → cada integrante → cierre.",
   size=11, color=MUTED, italic=True)

notes(s, "Agenda (10 s). Queda claro quién expone qué y en qué orden.")

# ================================================================ SLIDES POR INTEGRANTE (2 c/u)
def persona_slide(idx, eyebrow, title, sub, left_items, img_path, img_caption,
                  right_title=None, right_items=None, note=""):
    s = add_slide()
    header(s, eyebrow, title, sub)
    card(s, 0.75, 1.92, 6.30, 4.95, accent=BLUE, title=None, items=left_items, size=12.5)
    # imagen a la derecha
    if img_path:
        picture(s, img_path, 7.30, 2.00, 5.28, 3.30, border=True, radius=0.03)
        tb(s, 7.30, 5.38, 5.28, 0.30, img_caption, size=9.5, color=MUTED, align=PP_ALIGN.CENTER)
    if right_title:
        rect(s, 7.30, 5.78, 5.28, 1.06, fill=BLUE_LIGHT, line=LINE,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE, adjust=0.10)
        tb(s, 7.50, 5.88, 4.9, 0.3, right_title, size=11, color=NAVY, bold=True)
        if right_items:
            bullets(s, 7.50, 6.22, 4.95, 0.6, right_items, size=9.5, gap=4)
    footer(s, idx)
    notes(s, note)
    return s

# ---- I1-A Seguridad
persona_slide(
    7, "Integrante 1 · José Patricio Hidalgo Salas",
    "Seguridad: el token que abre todas las puertas",
    "identidad-service · /api/auth · login y roles",
    [
        "**Login:** usuario y contraseña van a identity-service; BCrypt compara el hash, nunca guarda la clave.",
        "**JWT:** devuelve un token firmado con sub, usuarioId, nombres, apellidos y rol.",
        "**Un solo secreto:** los 4 servicios validan el mismo JWT_SECRET, no consultan la base de usuarios.",
        "**Roles:** ADMIN (todo) y OPERADOR (caja, sin anular).",
        "**Sin permiso:** 401 sin token y 403 sin rol, igual que antes del cambio.",
    ],
    f"{IM}/01-panel.png",
    "Cabecera con «Administrador · ADMIN»: prueba de sesión iniciada (captura real).",
    right_title="En la demo",
    right_items=[
        "Login correcto e incorrecto.",
        "Identidad y rol visibles en la cabecera.",
    ],
    note="I1-A (45 s). Mostrar login, explicar BCrypt + JWT y los dos roles. Captura: sesión iniciada como administrador."
)

# ---- I1-B Gateway
persona_slide(
    8, "Integrante 1 · José Patricio Hidalgo Salas",
    "API Gateway: una sola puerta de entrada",
    "api-gateway · puerto 9101 · 10 rutas por prefijo",
    [
        "**Ruteo:** /api/auth/** → identity, /api/socios,giros,puestos,servicios,bancos → catalog, /api/cuentas,recibos,egresos,reportes → cash.",
        "**Sin Eureka:** las URLs vienen de variables de entorno; nada de descubrimiento de servicios.",
        "**Prefijo desconocido:** responde 404 del gateway (no 500 ni timeout).",
        "**CORS centralizado** para el frontend en :4400.",
        "**Salud:** /actuator/health y rutas visibles en /actuator/gateway/routes.",
    ],
    f"{IM}/01-panel.png",
    "Cabecera del frontend ya apunta al gateway; se ve \"Administrador · ADMIN\".",
    right_title="En la demo",
    right_items=[
        "curl al health de los 4 servicios.",
        "curl a /api/inexistente → 404 del gateway.",
    ],
    note="I1-B (45 s). Gateway como puerta única; mostrar los 10 rutas y un 404 intencional."
)

# ---- I2-A Catálogo
persona_slide(
    9, "Integrante 2 · Gerardo Matías Calizaya Alvarez",
    "catalog-service: los datos que todos necesitan",
    "catalog-service · puerto 9103 · 25 operaciones",
    [
        "**Cinco catálogos:** socios, giros, puestos, servicios y bancos.",
        "**Puesto:** giro obligatorio, socio opcional, vigencia con fechas de inicio y fin.",
        "**Servicio:** fijo o dependiente de consumo; moneda y si se cobra al puesto o al socio.",
        "**Mismas validaciones** y mensajes de error que el monolito.",
        "**Cero cambios de contrato:** el JSON de salida es idéntico al de antes.",
    ],
    f"{IM}/02-socios.png",
    "Catálogo de socios con 25 registros (captura real del front).",
    right_title="En la demo",
    right_items=[
        "CRUD de un puesto con socio opcional.",
        "Servicio fijo vs. servicio por consumo.",
    ],
    note="I2-A (45 s). Los cinco catálogos y las reglas de negocio que no cambiaron."
)

# ---- I2-B Fronteras
persona_slide(
    10, "Integrante 2 · Gerardo Matías Calizaya Alvarez",
    "Un dato, un solo dueño",
    "fronteras de datos · snapshots · sin FK cruzadas",
    [
        "**Cada tabla tiene un dueño:** solo su servicio la escribe.",
        "**Caja no tiene FK hacia catálogo:** guarda IDs más una copia (snapshot) congelada al emitir el documento.",
        "**Ejemplo:** un recibo guarda los nombres del socio tal como estaban al cobrar; si el socio cambia después, el recibo histórico no se altera.",
        "**Lecturas cruzadas por HTTP** con CatalogoClient, timeouts 1.5 s / 3 s y sin reintentos (evita duplicar recibos).",
    ],
    None, None,
    right_title="¿Por qué así?",
    right_items=[
        "Independencia real entre servicios.",
        "Trazabilidad de documentos sin depender de datos ajenos.",
    ],
    note="I2-B (45 s). Principio de un dueño por dato y snapshots; evita acoplar bases de datos."
)

# ---- I3-A Cobranza
persona_slide(
    11, "Integrante 3 · Juan Alonso Esteban Pajuelo Chipana",
    "Caja: de la deuda al pago en un solo flujo",
    "cash-service · puerto 9104 · cuentas y recibos",
    [
        "**Cuentas individuales o masivas** por puesto o socio, filtradas por etapa.",
        "**Cada cuenta:** servicio + periodo + monto (o costo unitario × consumo).",
        "**Pago o exoneración** individual y en lote, con total seleccionado visible.",
        "**Estados:** PENDIENTE → PAGADA / EXONERADA; una deuda no se paga dos veces.",
        "**Consulta** de cuentas por puesto o socio con resumen consolidado.",
    ],
    f"{IM}/03-cuentas.png",
    "Generación de cuentas y tabla de cobranza (captura real).",
    right_title="En la demo",
    right_items=[
        "Cuenta masiva y pago en lote con selección múltiple.",
        "Voucher detallado del recibo emitido.",
    ],
    note="I3-A (45 s). Flujo completo de cobranza: generar, seleccionar, pagar/exonerar y voucher."
)

# ---- I3-B Correlativos y recibos
persona_slide(
    12, "Integrante 3 · Juan Alonso Esteban Pajuelo Chipana",
    "Recibos: numeración sin huecos ni repetidos",
    "correlativos transaccionales · recibos y canjes",
    [
        "**Correlativo único:** REC-000221, bloqueado en su propia tabla dentro de la transacción.",
        "**Si falla la transacción, el número no se consume.**",
        "**Un recibo por pago** con destinatario, concepto y usuario que cobró.",
        "**Canjes bancarios e ingresos externos** usan el mismo mecanismo; simulados, no mueven dinero real.",
        "**0 duplicados** verificados en pruebas de correlativos.",
    ],
    f"{IM}/04-recibos.png",
    "Recibos de ingreso con correlativos REC-000221 en adelante (captura real).",
    right_title="En la demo",
    right_items=[
        "Emisión de recibo y apertura de su voucher.",
        "Canje bancario simulado (sin pasarela).",
    ],
    note="I3-B (45 s). Correlativos transaccionales y el alcance académico: no hay banco real."
)

# ---- I4-A Tesorería
persona_slide(
    13, "Integrante 4 · Salvador Roque Tupa Churata",
    "Tesorería: egresos y reportes",
    "cash-service · egresos, canjes y exportaciones",
    [
        "**Egresos individuales o por CSV**, con procesamiento y anulación.",
        "**Comprobantes correlativos** para cada egreso registrado.",
        "**Reportes XLSX:** diarios, mensuales, de socios, de egresos y de bancos.",
        "**Catálogo de bancos** referencial para canjes simulados.",
        "**Todo queda como constancia interna:** no mueve dinero ni se conecta a terceros.",
    ],
    f"{IM}/05-egresos.png",
    "Registro de egresos con historial y botón Procesar (captura real).",
    right_title="En la demo",
    right_items=[
        "Importar CSV de egresos y procesar.",
        "Descargar reporte XLSX mensual.",
    ],
    note="I4-A (45 s). Egresos, CSV y reportes XLSX; aclarar que el módulo bancario es simulación académica."
)

# ---- I4-B Frontend
persona_slide(
    14, "Integrante 4 · Salvador Roque Tupa Churata",
    "El frontend no se reescribió",
    "Angular 21 · único cambio: apiUrl al gateway",
    [
        "**Angular 21 con 13 pantallas y 10 servicios HTTP** se mantienen intactos.",
        "**Un solo cambio de configuración:** apiUrl de 9090 a 9101 (gateway).",
        "**Interceptor** agrega el JWT a cada petición.",
        "**Guard** protege las rutas y los roles ocultan acciones según el usuario.",
        "**Contrato JSON idéntico:** mismos nombres de campos y anidamientos que antes.",
    ],
    f"{IM}/06-reportes.png",
    "Pantalla de Reportes con exportaciones XLSX (captura real).",
    right_title="En la demo",
    right_items=[
        "Navegar por todas las pantallas sin recargar sesión.",
        "Cabecera muestra rol y permite cerrar sesión.",
    ],
    note="I4-B (45 s). Qué cambió en el frontend: una línea de configuración; nada de pantallas reescritas."
)

# ---- I5-A Migración
persona_slide(
    15, "Integrante 5 · Oscar Eduardo Carrillo Bautista",
    "Migrar sin apagar el sistema (strangler)",
    "patrón strangler · 4 estados · reversa por configuración",
    [
        "**Estado 0:** monolito puro en :9090.",
        "**Estado 1:** gateway + identidad; el resto sigue en el monolito.",
        "**Estado 2:** gateway + identidad + catálogo; caja aún en monolito.",
        "**Estado 3 (final):** gateway → identity, catalog y cash; el monolito queda como respaldo sin tráfico del frontend.",
        "**Cada fase deja el sistema funcionando** y es reversible solo cambiando configuración.",
    ],
    None, None,
    right_title="En el repo",
    right_items=[
        "80 commits en 17 ramas, cada tarea con su PR.",
        "Monolito intacto en backend/ como respaldo.",
    ],
    note="I5-A (45 s). Explicar el patrón strangler y los 4 estados; énfasis en que nunca se apagó el sistema."
)

# ---- I5-B Calidad
persona_slide(
    16, "Integrante 5 · Oscar Eduardo Carrillo Bautista",
    "Cómo sabemos que funciona",
    "calidad, contratos y operación",
    [
        "**52 operaciones documentadas** en OpenAPI/Swagger: 2 (identity) + 25 (catalog) + 25 (cash).",
        "**Colección Postman** con entornos: recorre login → catálogo → deuda → pago → recibo → reporte.",
        "**Pruebas JUnit/Mockito** por servicio y build de Angular.",
        "**Flyway:** 11 migraciones crean las 3 bases desde cero.",
        "**Actuator health** en los 4 servicios; gateway responde 404 a prefijos no listados.",
    ],
    f"{IM}/09-swagger.png",
    "Swagger de catalog-service (captura real de Chrome DevTools).",
    right_title="En la demo",
    right_items=[
        "Abrir /v3/api-docs de un servicio.",
        "Ejecutar la colección Postman completa.",
    ],
    note="I5-B (45 s). Evidencia de calidad: Swagger, Postman, pruebas y salud de los servicios."
)

# ================================================================ SLIDE 17 · RESULTADOS
s = add_slide()
header(s, "Resultados", "Lo que entregamos en números",
       sub="Cifras clave del proyecto de microservicios.")
kpis = [
    ("3 servicios + 1 gateway", "identity, catalog, cash y api-gateway", BLUE),
    ("52 endpoints", "documentados en OpenAPI/Swagger", GREEN),
    ("3 bases de datos", "una por servicio, creadas con Flyway", ORANGE),
    ("10 rutas", "del gateway, ninguna al monolito", VIOLET),
    ("2 roles", "ADMIN y OPERADOR con permisos intactos", NAVY),
    ("0 pasarela real", "bancos y pagos son simulación académica", RED),
]
kx = 0.75
for valor, etiqueta, ac in kpis:
    kpi(s, kx, 1.95, 1.88, 1.05, valor, etiqueta, accent=ac, vsize=19, lsize=9.5)
    kx += 1.98
card(s, 0.75, 3.35, 7.35, 3.15, accent=NAVY, title="Lecciones del equipo",
     items=[
         "**El contrato manda:** migrar sin romper el JSON evitó tocar el frontend.",
         "**Sin transacciones distribuidas:** snapshots + un dueño por dato bastaron.",
         "**Simplicidad operativa:** sin Eureka, Kafka, sagas ni contenedores fue una decisión, no una omisión.",
         "**Lo que falta:** despliegue en contenedores, observabilidad y pruebas de carga.",
     ], size=12)
card(s, 8.30, 3.35, 4.28, 3.15, accent=GREEN, title="Ámbito académico",
     items=[
         "Sin pasarela de pagos ni bancos reales.",
         "Sin OAuth2 ni proveedor de identidad externo.",
         "Monolito ejecutable como respaldo manual.",
     ], size=11.5)
footer(s, 17)
notes(s, "Resultados (25 s). Repasar los KPI rápidos y las lecciones; no entrar en detalle de cada una.")

# ================================================================ SLIDE 18 · CIERRE
s = add_slide(bg=NAVY)
rect(s, 0, 0, 13.333, 7.5, fill=NAVY)
rect(s, 0, 0, 13.333, 0.16, fill=BLUE)
tb(s, 0.9, 2.1, 8.5, 1.2, "¿Preguntas?", size=48, color=WHITE, bold=True)
tb(s, 0.9, 3.25, 8.5, 0.5, "Gracias por su atención.", size=20, color=RGBColor(0xCF, 0xDD, 0xEA))
rect(s, 0.92, 4.0, 2.0, 0.045, fill=BLUE)
tb(s, 0.9, 4.3, 8.5, 0.4, "Grupo 2 · DSW II · Ciclo 6 · 2026", size=13, color=RGBColor(0x9F, 0xC0, 0xFF), bold=True)
tb(s, 0.9, 4.72, 11.5, 1.0,
   "Hidalgo Salas, José Patricio · Calizaya Alvarez, Gerardo Matías · Pajuelo Chipana, Juan Alonso Esteban · Tupa Churata, Salvador Roque · Carrillo Bautista, Oscar Eduardo",
   size=12, color=RGBColor(0xCF, 0xDD, 0xEA))
tb(s, 0.9, 6.1, 8.5, 0.4, "Código y documentación: repositorio del proyecto (microservices/, frontend/, backend/, docs/).", size=10.5, color=RGBColor(0x9F, 0xC0, 0xFF))

OUT = "docs/presentacion/Presentacion_Microservicios_Grupo2.pptx"
prs.save(OUT)
print("OK:", OUT, "slides:", len(prs.slides.__iter__.__self__._sldIdLst))